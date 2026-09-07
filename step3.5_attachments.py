# -*- coding: utf-8 -*-
"""
Step 3.5 — تبدیل پیوست‌های متنی به Markdown برای RAG (AnythingLLM)

ورودی : Extracted_Sample/attachments/          (166 فایل باینری)
خروجی : Extracted_Sample_text/attachments_text/ (md متنی با هدر parent)

قواعد:
- xlsx/xlsm : هر شیت → یک جدول Markdown؛ سلول‌های merged پر می‌شوند؛
              فونت قرمز → برچسب [قرمز] (برای ستون Q کسری‌ها)
- xls       : xlrd (فرمت قدیمی Data Base Calibration)
- pdf       : pdfminer.six؛ اگر متن < آستانه → PDF اسکن → علامت «نیاز به OCR»
- docx/pptx : python-docx / python-pptx
- jpg/png/rar/dwg/نامعلوم : صرف‌نظر + ثبت در گزارش
- هر فایل md هدر parent دارد تا به ایمیل مادر وصل شود.
- شیت/سطرهای بیش از سقف → برش با یادداشت.
"""
import os
import glob
import re

BASE = os.path.dirname(os.path.abspath(__file__))
ATT_DIR = os.path.join(BASE, "Extracted_Sample", "attachments")
MD_DIR = os.path.join(BASE, "Extracted_Sample")
OUT_DIR = os.path.join(BASE, "Extracted_Sample_text", "attachments_text")

MAX_ROWS = 1000          # سقف سطر هر شیت
MAX_COLS = 40            # سقف ستون
PDF_MIN_CHARS_PER_PAGE = 120   # زیر این = PDF اسکن
PDF_MAX_PAGES = 60       # سقف صفحات PDF

RED_MIN = (150, 0, 0)    # آستانه تشخیص فونت قرمز

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".gif"}
SKIP_EXT = {".rar", ".zip", ".dwg", ".exe", ".msg"}


def find_parent_md(att_name):
    """14031117_005_9_xxx.xlsx → نام فایل md والد (14031117_005_*.md)"""
    m = re.match(r"^(\d{8}_\d{3})_", att_name)
    if not m:
        return ""
    prefix = m.group(1)
    hits = glob.glob(os.path.join(MD_DIR, prefix + "_*.md"))
    return os.path.basename(hits[0]) if hits else ""


def esc(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ").replace("|", "\\|")
    return s.strip()


def cell_font_red(cell):
    """تشخیص فونت قرمز سلول (rgb یا indexed)"""
    try:
        c = cell.font.color
        if c is None:
            return False
        if c.type == "rgb" and c.rgb and isinstance(c.rgb, str) and len(c.rgb) == 8:
            r, g, b = int(c.rgb[2:4], 16), int(c.rgb[4:6], 16), int(c.rgb[6:8], 16)
            return r >= RED_MIN[0] and g < 100 and b < 100
        if c.type == "indexed" and c.indexed == 10:  # استاندارد قرمز اکسل
            return True
    except Exception:
        pass
    return False


def md_table(rows):
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    out = []
    header = rows[0]
    out.append("| " + " | ".join(esc(header[i]) if i < len(header) else "" for i in range(ncol)) + " |")
    out.append("|" + "---|" * ncol)
    for r in rows[1:]:
        out.append("| " + " | ".join(esc(r[i]) if i < len(r) else "" for i in range(ncol)) + " |")
    return "\n".join(out)


def compact_rows(rows):
    """حذف ستون‌های خالی انتهایی + ادغام سطرهای تکراری متوالی (عنوان merged)"""
    if not rows:
        return rows
    ncol = max(len(r) for r in rows)
    norm = []
    for r in rows:
        r = list(r) + [""] * (ncol - len(r))
        cells = []
        for x in r:
            if x is None:
                x = ""
            elif isinstance(x, float):
                x = round(x, 4)
                if x == int(x):
                    x = int(x)
            cells.append(str(x).strip())
        norm.append(cells)
    # آخرین ستون دارای داده
    last = 0
    for r in norm:
        for i in range(ncol - 1, -1, -1):
            if r[i]:
                last = max(last, i + 1)
                break
    norm = [r[:last] for r in norm]
    # سطرهای تکراری متوالی (نتیجه merge سطری عنوان)
    out = []
    for r in norm:
        if out and r == out[-1]:
            continue
        out.append(r)
    return out


def xlsx_to_md(path, att_name):
    """اکسل مدرن: هر شیت → جدول md؛ خروجی: (sections_md, meta)
    فایل حجیم → حالت سریع (read_only، بدون رنگ/merged)"""
    from openpyxl import load_workbook
    big = os.path.getsize(path) > 2 * 1024 * 1024
    wb = load_workbook(path, data_only=True, read_only=big)
    parts = []
    meta = []
    note_big = "" if not big else " (فایل حجیم — بدون تشخیص رنگ و سلول ادغامی)"
    for ws in wb.worksheets:
        merged_map = {}
        if not big:
            for rng in ws.merged_cells.ranges:
                anchor = ws.cell(rng.min_row, rng.min_col).value
                # merge افقی: فقط سلول اول مقدار دارد (بدون تکرار)
                # merge عمودی: همه سطرها مقدار لنگر را می‌گیرند
                if anchor is None:
                    continue
                horizontal = rng.max_row == rng.min_row
                for row in range(rng.min_row, rng.max_row + 1):
                    for col in range(rng.min_col, rng.max_col + 1):
                        if horizontal and col > rng.min_col:
                            merged_map[(row, col)] = ""
                        else:
                            merged_map[(row, col)] = anchor

        rows = []
        red_count = 0
        total_rows = ws.max_row or 0
        for r_idx, row in enumerate(
            ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS), 1
        ):
            vals = []
            has_data = False
            for c_idx, cell in enumerate(row, 1):
                if big:
                    v = cell.value
                else:
                    v = merged_map.get((r_idx, c_idx), cell.value)
                if v is not None and str(v).strip() != "":
                    has_data = True
                if not big and cell_font_red(cell) and v not in (None, ""):
                    red_count += 1
                    vals.append("[قرمز] " + str(v))
                else:
                    vals.append(v)
            if has_data:
                rows.append(vals)
        truncated = total_rows > MAX_ROWS

        if not rows:
            continue
        notes = []
        rows = compact_rows(rows)
        sec = "## شیت: %s\n\n" % esc(ws.title)
        sec += md_table(rows[:MAX_ROWS])
        if truncated:
            notes.append("(جدول بریده شد: %d سطر از %d)" % (MAX_ROWS, total_rows))
        if red_count:
            notes.append("(سلول‌های قرمز = کسری/هشدار: %d مورد، با برچسب [قرمز])" % red_count)
        if note_big:
            notes.append(note_big.strip(" ("))
        if notes:
            sec += "\n\n_" + "؛ ".join(notes) + "_"
        parts.append(sec)
        meta.append("%s (%d سطر)" % (ws.title, len(rows)))
    wb.close()
    return "\n\n".join(parts), meta


def xls_to_md(path, att_name):
    """اکسل قدیمی xls با xlrd"""
    import xlrd
    wb = xlrd.open_workbook(path)
    parts, meta = [], []
    for ws in wb.sheets():
        rows = []
        for r in range(min(ws.nrows, MAX_ROWS)):
            vals = []
            for c in range(min(ws.ncols, MAX_COLS)):
                v = ws.cell_value(r, c)
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                vals.append("" if v == "" else v)
            if any(str(x).strip() for x in vals):
                rows.append(vals)
        if not rows:
            continue
        rows = compact_rows(rows)
        sec = "## شیت: %s\n\n" % esc(ws.name)
        sec += md_table(rows)
        if ws.nrows > MAX_ROWS:
            sec += "\n\n_(جدول بریده شد: %d سطر از %d)_" % (MAX_ROWS, ws.nrows)
        parts.append(sec)
        meta.append("%s (%d سطر)" % (ws.name, len(rows)))
    return "\n\n".join(parts), meta


def pdf_to_md(path, att_name):
    """PDF: متن استخراج؛ اگر اسکن → علامت"""
    from pdfminer.high_level import extract_text
    try:
        text = extract_text(path, maxpages=PDF_MAX_PAGES)
    except Exception as e:
        return "", ["خطا: %s" % str(e)[:80]], "pdf-خطا"
    clean = re.sub(r"\n{3,}", "\n\n", text).strip()
    pages_hint = clean.count("\x0c") + 1
    if len(clean) < PDF_MIN_CHARS_PER_PAGE * max(pages_hint, 1):
        return "", ["%d صفحه — بدون متن قابل استخراج" % pages_hint], "pdf-اسکن"
    return clean, ["%d صفحه، %d کاراکتر متن" % (pages_hint, len(clean))], "pdf-متنی"


def docx_to_md(path, att_name):
    import docx
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower()
            if "heading" in style:
                parts.append("## " + p.text.strip())
            else:
                parts.append(p.text.strip())
    for t in d.tables:
        rows = [[c.text for c in row.cells] for row in t.rows]
        parts.append(md_table(rows))
    return "\n\n".join(parts), ["%d پاراگراف، %d جدول" % (len(d.paragraphs), len(d.tables))]


def pptx_to_md(path, att_name):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        if texts:
            parts.append("## اسلاید %d\n\n%s" % (i, "\n\n".join(texts)))
    return "\n\n".join(parts), ["%d اسلاید" % len(prs.slides)]


def write_att_md(out_name, att_name, kind, parent, notes, body):
    fm = [
        "---",
        "پیوست: %s" % att_name,
        "نوع: %s" % kind,
        "والد: %s" % (parent or "نامشخص"),
    ]
    for n in notes:
        fm.append("یادداشت: %s" % n)
    fm.append("---")
    content = "\n".join(fm) + "\n\n" + (body.strip() if body.strip() else "_(بدون محتوای متنی)_")
    with open(os.path.join(OUT_DIR, out_name), "w", encoding="utf-8") as f:
        f.write(content)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    files = sorted(glob.glob(os.path.join(ATT_DIR, "*.*")))
    report = []
    stats = {"md": 0, "skip": 0, "scan": 0, "err": 0}

    for fp in files:
        att_name = os.path.basename(fp)
        ext = os.path.splitext(att_name)[1].lower()
        stem = os.path.splitext(att_name)[0]
        parent = find_parent_md(att_name)
        print("processing:", att_name.encode("unicode_escape").decode()[:50], flush=True)

        try:
            if ext in IMAGE_EXT or ext in SKIP_EXT:
                stats["skip"] += 1
                report.append("SKIP  %-8s %s" % (ext, att_name))
                continue
            if ext == ".16":
                stats["skip"] += 1
                report.append("SKIP  bad-ext %s" % att_name)
                continue

            if ext in (".xlsx", ".xlsm"):
                body, meta = xlsx_to_md(fp, att_name)
                write_att_md(stem + ".md", att_name, "xlsx — " + "؛ ".join(meta), parent, meta, body)
            elif ext == ".xls":
                body, meta = xls_to_md(fp, att_name)
                write_att_md(stem + ".md", att_name, "xls — " + "؛ ".join(meta), parent, meta, body)
            elif ext == ".pdf":
                body, meta, kind = pdf_to_md(fp, att_name)
                if kind == "pdf-اسکن":
                    stats["scan"] += 1
                    write_att_md(stem + ".md", att_name, kind, parent, meta,
                                 "این PDF اسکن/تصویر است و متن قابل استخراج ندارد — نیاز به OCR (فاز ۲).")
                    report.append("SCAN  %s" % att_name)
                    continue
                write_att_md(stem + ".md", att_name, kind, parent,
                             meta + ["متن PDF (به‌ویژه فارسی RTL) ممکن است به‌هم‌ریخته باشد"], body)
            elif ext == ".docx":
                body, meta = docx_to_md(fp, att_name)
                write_att_md(stem + ".md", att_name, "docx", parent, meta, body)
            elif ext == ".pptx":
                body, meta = pptx_to_md(fp, att_name)
                write_att_md(stem + ".md", att_name, "pptx", parent, meta, body)
            else:
                stats["skip"] += 1
                report.append("SKIP  %-8s %s" % (ext, att_name))
                continue

            stats["md"] += 1
            report.append("OK    %s" % att_name)
        except Exception as e:
            stats["err"] += 1
            report.append("ERR   %s — %s" % (att_name, str(e)[:90]))

    with open(os.path.join(BASE, "Extracted_Sample_text", "attachments_report.txt"),
              "w", encoding="utf-8") as f:
        f.write("تبدیل‌شده: %d | اسکن (نیاز OCR): %d | صرف‌نظر: %d | خطا: %d\n\n" % (
            stats["md"], stats["scan"], stats["skip"], stats["err"]))
        f.write("\n".join(report))
    print("done", stats)


if __name__ == "__main__":
    main()
