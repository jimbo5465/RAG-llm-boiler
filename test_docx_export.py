# -*- coding: utf-8 -*-
import os
import io
import google_drive_ocr as g_ocr
from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload
from docx import Document
from project_paths import TEST_INPUT_DIR, TEST_OUTPUT_DIR

drive_service = g_ocr.get_drive_service()

pdf_path = os.path.join(TEST_INPUT_DIR, "36423469 (1).pdf")
file_metadata = {
    'name': 'test_docs_docx_export',
    'mimeType': 'application/vnd.google-apps.document'
}
media = MediaFileUpload(pdf_path, mimetype='application/pdf', resumable=True)
uploaded = drive_service.files().create(body=file_metadata, media_body=media, fields='id').execute()
doc_id = uploaded.get('id')
print("Uploaded Doc ID:", doc_id)

os.makedirs(TEST_OUTPUT_DIR, exist_ok=True)
docx_path = os.path.join(TEST_OUTPUT_DIR, "exported_doc.docx")

try:
    # صادرات به فرمت DOCX به جای Markdown
    request = drive_service.files().export_media(
        fileId=doc_id,
        mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    with io.FileIO(docx_path, 'wb') as fh:
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
    print("Downloaded DOCX successfully!")

    # بررسی ساختار جداول در فایل DOCX
    doc = Document(docx_path)
    print(f"Total tables in DOCX: {len(doc.tables)}")
    print(f"Total paragraphs in DOCX: {len(doc.paragraphs)}")

    for t_idx, table in enumerate(doc.tables):
        print(f"\n--- DOCX Table {t_idx+1}: {len(table.rows)} rows x {len(table.columns)} cols ---")
        for r_idx, row in enumerate(table.rows[:6]):
            row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"  Row {r_idx}: {row_vals}")

finally:
    drive_service.files().delete(fileId=doc_id).execute()
    print("Cleaned up Doc from Google Drive.")
