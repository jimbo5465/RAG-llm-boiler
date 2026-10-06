# -*- coding: utf-8 -*-
"""
ماژول استخراج محلی پیوست‌ها:
- بازگشایی خودکار فایل‌های ZIP
- تبدیل اسناد آفیس (Excel: xlsx/xls, Word: docx, PowerPoint: pptx) به Markdown ساختاریافته
- فیلتر و ارجاع PDFها و تصاویر به ماژول Google OCR
"""
from project_paths import PROJECT_ROOT
import os
import glob
import re
import zipfile
import hashlib
from pathlib import Path
import shutil
import uuid

BASE = PROJECT_ROOT
ATT_DIR = os.path.join(BASE, "Extracted_Sample", "attachments")
MD_DIR = os.path.join(BASE, "Extracted_Sample")
OUT_DIR = os.path.join(BASE, "Extracted_Sample_text", "attachments_text")

MAX_ROWS = 1000          # سقف سطر هر شیت
MAX_COLS = 40            # سقف ستون
RED_MIN = (150, 0, 0)    # آستانه تشخیص فونت قرمز

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".gif"}
SKIP_EXT = {".rar", ".dwg", ".exe", ".msg"}


def esc(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ").replace("|", "\\|")
    return s.strip()


def cell_font_red(cell):
    """تشخیص فونت قرمز سلول (rgb یا indexed)"""
    try:
        c = cell.font.color
        if c is None:
            return False
        if c.type == "rgb" and c.rgb:
            r = int(str(c.rgb)[0:2], 16)
            g = int(str(c.rgb)[2:4], 16)
            b = int(str(c.rgb)[4:6], 16)
            return (r, g, b) >= RED_MIN and r > g + 40 and r > b + 40
        if c.type == "indexed" and c.indexed in (2, 10, 16):
            return True
    except Exception:
        pass
    return False


def md_table(rows):
    """لیست سطور (هر سطر لیستی از رشته‌ها) → جدول Markdown تمیز"""
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    if ncols == 0:
        return ""
    norm = []
    for r in rows:
        padded = [esc(x) for x in r] + [""] * (ncols - len(r))
        norm.append(padded)
    header = norm[0]
    sep = ["---"] * ncols
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(sep) + " |",
    ]
    for r in norm[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def compact_rows(rows):
    """حذف ستون‌ها و سطرهای کاملاً خالی"""
    if not rows:
        return []
    rows = [r for r in rows if any(str(x).strip() for x in r)]
    if not rows:
        return []
    ncols = max(len(r) for r in rows)
    non_empty_cols = []
    for c in range(ncols):
        col_vals = [r[c] for r in rows if c < len(r)]
        if any(str(x).strip() for x in col_vals):
            non_empty_cols.append(c)
    compact = []
    for r in rows:
        compact.append([r[c] if c < len(r) else "" for c in non_empty_cols])
    return compact


def xlsx_to_md(path, att_name):
    """اکسل جدید با openpyxl — تشخیص فونت قرمز و خلاصه‌سازی هوشمند"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    parts = []
    meta = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        red_cells = []
        for r_idx, row in enumerate(ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS), 1):
            row_vals = []
            for c_idx, cell in enumerate(row, 1):
                v = cell.value
                if v is None:
                    row_vals.append("")
                else:
                    if isinstance(v, float) and v.is_integer():
                        v = int(v)
                    row_vals.append(v)
                    if cell_font_red(cell):
                        red_cells.append("%s (سطر %d ستون %d)" % (str(v)[:40], r_idx, c_idx))
            if any(str(x).strip() for x in row_vals):
                rows.append(row_vals)
        if not rows:
            continue
        rows = compact_rows(rows)
        sec = "## شیت: %s\n\n" % esc(sheet_name)
        sec += md_table(rows)
        notes = []
        if ws.max_row and ws.max_row > MAX_ROWS:
            notes.append("جدول بریده شد: %d سطر از %d" % (MAX_ROWS, ws.max_row))
        if red_cells:
            note_red = "سلول‌های با فونت قرمز (موارد با اهمیت بالا/عدم انطباق/اقلام بازمانده): "
            note_red += "؛ ".join(red_cells[:15])
            if len(red_cells) > 15:
                note_red += " (... و %d مورد دیگر)" % (len(red_cells) - 15)
            notes.append(note_red)
        if notes:
            sec += "\n\n_" + "؛ ".join(notes) + "_"
        parts.append(sec)
        meta.append("%s (%d سطر)" % (sheet_name, len(rows)))
    wb.close()
    return "\n\n".join(parts), meta


def xls_to_md(path, att_name):
    """اکسل قدیمی xls با xlrd"""
    import xlrd
    wb = xlrd.open_workbook(path)
    parts, meta = [], []
    for ws in wb.sheets():
        rows = []
        for r in range(min(ws.nrows, MAX_ROWS)):
            vals = []
            for c in range(min(ws.ncols, MAX_COLS)):
                v = ws.cell_value(r, c)
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                vals.append("" if v == "" else v)
            if any(str(x).strip() for x in vals):
                rows.append(vals)
        if not rows:
            continue
        rows = compact_rows(rows)
        sec = "## شیت: %s\n\n" % esc(ws.name)
        sec += md_table(rows)
        if ws.nrows > MAX_ROWS:
            sec += "\n\n_(جدول بریده شد: %d سطر از %d)_" % (MAX_ROWS, ws.nrows)
        parts.append(sec)
        meta.append("%s (%d سطر)" % (ws.name, len(rows)))
    return "\n\n".join(parts), meta


def docx_to_md(path, att_name):
    """تبدیل Word docx به Markdown"""
    import docx
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower()
            if "heading" in style:
                parts.append("## " + p.text.strip())
            else:
                parts.append(p.text.strip())
    for t in d.tables:
        rows = [[c.text for c in row.cells] for row in t.rows]
        parts.append(md_table(rows))
    return "\n\n".join(parts), ["%d پاراگراف، %d جدول" % (len(d.paragraphs), len(d.tables))]


def pptx_to_md(path, att_name):
    """تبدیل PowerPoint pptx به Markdown"""
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        if texts:
            parts.append("## اسلاید %d\n\n%s" % (i, "\n\n".join(texts)))
    return "\n\n".join(parts), ["%d اسلاید" % len(prs.slides)]


def pdf_digital_to_md(path, att_name):
    """استخراج آفلاین و محلی متن از PDFهای دیجیتال بدون نیاز به اینترنت و OCR"""
    import pymupdf
    import unicodedata

    try:
        doc = pymupdf.open(path)
    except Exception:
        return None, []

    total_pages = len(doc)
    page_texts = []
    total_chars = 0

    for p_no, page in enumerate(doc, 1):
        text = page.get_text("text")
        if not text or not text.strip():
            continue
        clean_p = unicodedata.normalize("NFC", text).strip()
        lines = [l.strip() for l in clean_p.splitlines() if l.strip()]
        if lines:
            page_content = "\n".join(lines)
            total_chars += len(page_content)
            if total_pages > 1:
                page_texts.append(f"## صفحه {p_no}\n\n{page_content}")
            else:
                page_texts.append(page_content)
    doc.close()

    # اگر متن بسیار کم بود یا سند اسکن/تصویر بود، نادیده گرفته شود تا به OCR سپرده شود
    if total_chars < 80:
        return None, []

    return "\n\n---\n\n".join(page_texts), [f"{total_pages} صفحه", f"{total_chars} کاراکتر"]


import json

def write_att_md(out_dir, out_name, att_name, kind, parents, notes, body):
    fm = [
        "---",
        "پیوست: %s" % att_name,
        "نوع: %s" % kind,
    ]
    if isinstance(parents, list) and parents:
        fm.append("ایمیل‌های مرتبط (والد):")
        for p in parents:
            fm.append("  - %s" % p)
    elif isinstance(parents, str) and parents.strip():
        fm.append("ایمیل‌های مرتبط (والد):")
        fm.append("  - %s" % parents.strip())
    else:
        fm.append("ایمیل‌های مرتبط (والد): نامشخص")

    for n in notes:
        fm.append("یادداشت: %s" % n)
    fm.append("---")
    content = "\n".join(fm) + "\n\n" + (body.strip() if body.strip() else "_(بدون محتوای متنی)_")
    with open(os.path.join(out_dir, out_name), "w", encoding="utf-8") as f:
        f.write(content)


def fit_archive_filename(directory, name):
    limit = max(60, 240 - len(str(Path(directory).resolve())) - 1)
    if len(name) <= limit:
        return name
    suffix = Path(name).suffix
    return name[:limit - len(suffix) - 13] + '_' + hashlib.sha256(name.encode('utf-8')).hexdigest()[:12] + suffix


def open_rar(path):
    try:
        import rarfile
    except ImportError as error:
        raise RuntimeError('برای بازکردن RAR کتابخانه rarfile لازم است: py -3 -m pip install rarfile') from error
    for base in (os.environ.get('ProgramFiles'), os.environ.get('ProgramFiles(x86)')):
        if base:
            executable = Path(base) / 'WinRAR' / 'UnRAR.exe'
            if executable.is_file():
                rarfile.UNRAR_TOOL = str(executable)
                break
    rarfile.tool_setup()
    return rarfile.RarFile(str(path))


def extract_and_flatten_archives(att_dir, attachment_manifest=None, archive_paths=None):
    """بازکردن ZIP/RAR؛ حفظ اصل آرشیو، نام‌های یکتا و همه ایمیل‌های والد."""
    errors = 0
    archives = sorted(p for p in (map(Path, archive_paths) if archive_paths is not None else Path(att_dir).iterdir())
                      if p.is_file() and p.suffix.lower() in ('.zip', '.rar'))
    for archive_path in archives:
        prefix = 'unzipped' if archive_path.suffix.lower() == '.zip' else 'unrar'
        try:
            opener = zipfile.ZipFile if archive_path.suffix.lower() == '.zip' else open_rar
            with opener(archive_path) as archive:
                for member in archive.infolist():
                    if member.is_dir():
                        continue
                    if hasattr(member, 'is_symlink') and member.is_symlink():
                        continue
                    filename = member.filename.replace('\\', '/').rsplit('/', 1)[-1]
                    if not filename or filename.startswith(".") or filename.startswith("__MACOSX"):
                        continue
                    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', filename).rstrip(' .')
                    if not filename:
                        continue
                    target = Path(att_dir) / fit_archive_filename(att_dir, f'{archive_path.stem}_{prefix}_{filename}')
                    temporary = Path(att_dir) / f'_archive_{uuid.uuid4().hex}.tmp'
                    try:
                        with archive.open(member) as source, temporary.open('wb') as output:
                            shutil.copyfileobj(source, output)
                        data_hash = hashlib.sha256(temporary.read_bytes()).hexdigest()
                        if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest() != data_hash:
                            target = target.with_name(f'{target.stem}_{data_hash[:12]}{target.suffix}')
                        if not target.exists():
                            temporary.replace(target)
                        if attachment_manifest is not None:
                            parents = attachment_manifest.setdefault(target.name, [])
                            for parent in attachment_manifest.get(archive_path.name, []):
                                if parent not in parents:
                                    parents.append(parent)
                        print(f'[+] استخراج از {archive_path.name}: {target.name}')
                    finally:
                        temporary.unlink(missing_ok=True)
        except Exception as e:
            errors += 1
            print(f"[-] خطا در بازگشایی آرشیو {archive_path.name}: {e}")
    return errors


def extract_and_flatten_zips(att_dir):
    """نام قدیمی برای سازگاری؛ اکنون ZIP و RAR هر دو باز می‌شوند."""
    return extract_and_flatten_archives(att_dir)


def process_attachments(att_dir=ATT_DIR, md_dir=MD_DIR, out_dir=OUT_DIR, attachment_manifest=None,
                        selected_files=None):
    """پردازش آفلاین اسناد اکسل، ورد، پاورپوینت و بازگشایی زیپ‌ها با پشتیبانی از مانیفست والدین"""
    os.makedirs(out_dir, exist_ok=True)
    
    if attachment_manifest is None:
        manifest_path = os.path.join(os.path.dirname(att_dir), "attachment_manifest.json")
        if not os.path.exists(manifest_path):
            manifest_path = os.path.join(att_dir, "attachment_manifest.json")
        if os.path.exists(manifest_path):
            try:
                with open(manifest_path, "r", encoding="utf-8") as mf:
                    attachment_manifest = json.load(mf)
            except Exception:
                attachment_manifest = {}
        else:
            attachment_manifest = {}

    # ۱. بازگشایی خودکار فایل‌های ZIP
    archive_errors = extract_and_flatten_archives(att_dir, attachment_manifest) if selected_files is None else 0
    
    files = sorted(glob.glob(os.path.join(att_dir, "*.*"))) if selected_files is None else sorted(map(str, selected_files))
    report = []
    stats = {"md": 0, "skip": 0, "err": archive_errors}

    for fp in files:
        att_name = os.path.basename(fp)
        ext = os.path.splitext(att_name)[1].lower()
        stem = os.path.splitext(att_name)[0]
        
        # فایل‌های تصویر و MSG توسط سایر ماژول‌ها پردازش می‌شوند
        if ext in (".jpg", ".jpeg", ".png", ".bmp", ".msg") or ext in SKIP_EXT:
            continue

        # پیدا کردن والدین از attachment_manifest یا md_dir
        parents = attachment_manifest.get(att_name, [])
        if not parents:
            m = re.match(r"^(\d{8}_\d{3})_", att_name)
            if m:
                prefix = m.group(1)
                hits = glob.glob(os.path.join(md_dir, prefix + "_*.md"))
                if hits:
                    parents = [os.path.basename(hits[0])]

        try:
            if ext == ".pdf":
                # PDF در مسیر ترکیبی صفحه‌به‌صفحه پردازش می‌شود؛ استخراج محلی بدون
                # کنترل کیفیت قدیمی نباید این مرحله را دور بزند.
                stats["skip"] += 1
                report.append("GOCR  %s" % att_name)
                continue
            elif ext in (".xlsx", ".xlsm"):
                body, meta = xlsx_to_md(fp, att_name)
                write_att_md(out_dir, stem + ".md", att_name, "xlsx — " + "؛ ".join(meta), parents, meta, body)
            elif ext == ".xls":
                body, meta = xls_to_md(fp, att_name)
                write_att_md(out_dir, stem + ".md", att_name, "xls — " + "؛ ".join(meta), parents, meta, body)
            elif ext == ".docx":
                body, meta = docx_to_md(fp, att_name)
                write_att_md(out_dir, stem + ".md", att_name, "docx", parents, meta, body)
            elif ext == ".pptx":
                body, meta = pptx_to_md(fp, att_name)
                write_att_md(out_dir, stem + ".md", att_name, "pptx", parents, meta, body)
            else:
                stats["skip"] += 1
                continue

            stats["md"] += 1
            report.append("OK    %s" % att_name)
        except Exception as e:
            stats["err"] += 1
            report.append("ERR   %s: %s" % (att_name, str(e)[:60]))

    return stats
